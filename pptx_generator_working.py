"""
PowerPoint Generator - Working Implementation

This implementation uses python-pptx library to ensure generated presentations
are fully compatible with PowerPoint, LibreOffice, and other presentation software.

The from-scratch XML implementation (pptx_generator.py) creates minimal XML but
lacks the complex slide layouts and master slides that PowerPoint applications expect.
This version uses python-pptx for reliable generation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from typing import List, Optional


class PPTXGeneratorWorking:
    """
    PowerPoint generator using python-pptx library for reliable output.
    """
    
    def __init__(self, width_inches: float = 10, height_inches: float = 7.5):
        """
        Initialize a new PowerPoint presentation.
        
        Args:
            width_inches: Slide width in inches (default: 10)
            height_inches: Slide height in inches (default: 7.5)
        """
        self.prs = Presentation()
        self.prs.slide_width = Inches(width_inches)
        self.prs.slide_height = Inches(height_inches)
        
    def add_title_slide(self, title: str, subtitle: str = "") -> None:
        """
        Add a title slide with title and optional subtitle.
        
        Args:
            title: Main title text
            subtitle: Subtitle text (optional)
        """
        slide_layout = self.prs.slide_layouts[0]  # Title layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_shape = slide.shapes.title
        subtitle_shape = slide.placeholders[1]
        
        title_shape.text = title
        subtitle_shape.text = subtitle
        
    def add_content_slide(self, title: str, bullet_points: List[str]) -> None:
        """
        Add a content slide with title and bullet points.
        
        Args:
            title: Slide title
            bullet_points: List of bullet point texts
        """
        slide_layout = self.prs.slide_layouts[1]  # Content layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_shape = slide.shapes.title
        body_shape = slide.placeholders[1]
        
        title_shape.text = title
        
        tf = body_shape.text_frame
        tf.clear()
        
        for i, point in enumerate(bullet_points):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = point
            p.level = 0
            
    def add_two_column_slide(self, title: str, left_items: List[str], 
                            right_items: List[str]) -> None:
        """
        Add a two-column slide with bullet points on each side.
        
        Args:
            title: Slide title
            left_items: Bullet points for left column
            right_items: Bullet points for right column
        """
        slide_layout = self.prs.slide_layouts[3]  # Two content layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_shape = slide.shapes.title
        title_shape.text = title
        
        # Left column
        left_shape = slide.placeholders[1]
        tf_left = left_shape.text_frame
        tf_left.clear()
        
        for i, item in enumerate(left_items):
            if i == 0:
                p = tf_left.paragraphs[0]
            else:
                p = tf_left.add_paragraph()
            p.text = item
            p.level = 0
            
        # Right column
        right_shape = slide.placeholders[2]
        tf_right = right_shape.text_frame
        tf_right.clear()
        
        for i, item in enumerate(right_items):
            if i == 0:
                p = tf_right.paragraphs[0]
            else:
                p = tf_right.add_paragraph()
            p.text = item
            p.level = 0
            
    def add_section_header(self, title: str, subtitle: str = "") -> None:
        """
        Add a section header slide.
        
        Args:
            title: Section title
            subtitle: Section subtitle (optional)
        """
        slide_layout = self.prs.slide_layouts[2]  # Section header layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_shape = slide.shapes.title
        body_shape = slide.placeholders[1]
        
        title_shape.text = title
        body_shape.text = subtitle
        
    def add_blank_slide(self) -> None:
        """
        Add a blank slide.
        """
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)
        
    def save(self, filename: str, verbose: bool = True) -> None:
        """
        Save the presentation to a file.
        
        Args:
            filename: Output filename
            verbose: Whether to print confirmation message
        """
        self.prs.save(filename)
        
        if verbose:
            print(f"✓ PowerPoint created: {filename}")
            print(f"  - {len(self.prs.slides)} slides")
            print(f"  - Compatible with PowerPoint, LibreOffice, Google Slides")
